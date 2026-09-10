import json, os
from pathlib import Path
import streamlit as st
from ai.quiz_generator import generate_gemini, generate_groq

BASE=Path(__file__).parent
plan=json.loads((BASE/"data/plan_english_translation_2026.json").read_text(encoding="utf-8"))
COURSES=plan["courses"]; COURSES_DIR=BASE/"courses"
PRIMARY="#005980"; DARK="#003A54"; GOLD="#C59B27"
st.set_page_config(page_title="منصة التميز الأكاديمي | جامعة الملك سعود",page_icon="🎓",layout="wide")
st.markdown(f"""<style>
@import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;800&display=swap');
html,body,[class*="css"]{{font-family:Cairo,sans-serif;direction:rtl}} .stApp{{background:#f5f7fa}}
section[data-testid="stSidebar"]{{background:linear-gradient(180deg,{DARK},{PRIMARY})}} section[data-testid="stSidebar"] *{{color:white!important}}
.top{{background:linear-gradient(135deg,{DARK},{PRIMARY});color:white;padding:34px 40px;border-radius:0 0 24px 24px;margin:-1rem -1rem 25px}}
.top h1{{margin:0;font-size:34px;font-weight:800}} .hero,.card{{background:white;border:1px solid #e5e7eb;border-radius:20px;padding:24px;box-shadow:0 5px 22px #00000008}}
.title{{font-size:30px;font-weight:800;color:{DARK}}}.badge{{background:#eaf5f8;color:{PRIMARY};padding:5px 10px;border-radius:999px;font-size:12px;font-weight:700}}
.metric{{font-size:28px;font-weight:800;color:{PRIMARY}}}.q{{background:white;border-right:5px solid {PRIMARY};padding:20px;border-radius:15px}}
.ok{{background:#ecfdf5;border:1px solid #86efac;padding:16px;border-radius:15px;color:#166534}}
.no{{background:#fef2f2;border:1px solid #fca5a5;padding:16px;border-radius:15px;color:#991b1b}}
</style>""",unsafe_allow_html=True)

def folder(c): p=COURSES_DIR/c["id"]/"files"; p.mkdir(parents=True,exist_ok=True); return p
def extract(p):
    try:
        if p.suffix.lower() in [".txt",".md"]: return p.read_text(encoding="utf-8",errors="ignore")
        if p.suffix.lower()==".pdf":
            from pypdf import PdfReader; return "\n".join(x.extract_text() or "" for x in PdfReader(str(p)).pages)
        if p.suffix.lower()==".docx":
            from docx import Document; return "\n".join(x.text for x in Document(str(p)).paragraphs)
        if p.suffix.lower()==".pptx":
            from pptx import Presentation; return "\n".join(sh.text for s in Presentation(str(p)).slides for sh in s.shapes if hasattr(sh,"text"))
    except Exception: return ""
    return ""
def fallback(c,n):
    base=[("ما أفضل طريقة لدراسة المقرر؟",["الحفظ فقط","الفهم والتطبيق والاختبار الذاتي","التأجيل","التخمين"],"الفهم والتطبيق والاختبار الذاتي"),
    ("ما الأفضل عند وجود خطأ متكرر؟",["تجاهله","تسجيل سببه ومراجعته","حذفه","تغيير الموضوع"],"تسجيل سببه ومراجعته")]
    qs=[]
    for i in range(n):
        q,o,a=base[i%2]; qs.append({"question":q+" "+c["name"],"options":o,"answer":a,"explanation":"التعلم النشط وتحليل الأخطاء يساعدان على تثبيت الفهم."})
    return qs
def make_quiz(c,provider,key,n):
    texts=[extract(p) for p in folder(c).glob("*") if p.is_file()]
    context="\n".join(x for x in texts if x.strip())
    if not key or not context: return fallback(c,n),"تم استخدام بنك الأسئلة الاحتياطي."
    try:
        qs=generate_gemini(key,context,n) if provider=="Gemini" else generate_groq(key,context,n)
        return qs[:n],"تم إنشاء الاختبار من ملفات المقرر."
    except Exception as e: return fallback(c,n),f"تعذر استخدام الذكاء الاصطناعي؛ تم التحويل تلقائيًا للاحتياطي."

if "page" not in st.session_state: st.session_state.page="home"
if "cid" not in st.session_state: st.session_state.cid=None
if "quiz" not in st.session_state: st.session_state.quiz=None
if "i" not in st.session_state: st.session_state.i=0
if "score" not in st.session_state: st.session_state.score=0
if "answered" not in st.session_state: st.session_state.answered=False

with st.sidebar:
    st.markdown("<div style='text-align:center;font-size:48px'>🎓</div><h2 style='text-align:center'>منصة التميز الأكاديمي</h2><div style='text-align:center'>نموذج عرض مقترح</div>",unsafe_allow_html=True)
    if st.button("⌂ الرئيسية",use_container_width=True): st.session_state.page="home";st.rerun()
    with st.expander("🏛️ الكليات",True):
        if st.button("كلية اللغات وعلومها",use_container_width=True): st.session_state.page="college";st.rerun()
        if st.button("اللغة الإنجليزية والترجمة",use_container_width=True): st.session_state.page="program";st.rerun()
    provider=st.selectbox("مزود الذكاء الاصطناعي",["Gemini","Groq"])
    key=st.text_input("مفتاح API اختياري",type="password")
    st.caption("يمكن تشغيل العرض بدون مفتاح باستخدام بنك احتياطي.")
st.markdown(f'<div class="top"><h1>منصة التميز الأكاديمي</h1><div>بوابة أكاديمية مقترحة لطلاب كلية اللغات وعلومها — جامعة الملك سعود</div></div>',unsafe_allow_html=True)

if st.session_state.page=="home":
    st.markdown('<div class="hero"><div class="title">مرحبًا بك في المنصة الأكاديمية</div><p>الوصول إلى البرامج والمقررات والملفات والاختبارات الذكية في تجربة واحدة.</p></div>',unsafe_allow_html=True)
    a,b,c=st.columns(3)
    for col,t,v in [(a,"المقررات",len(COURSES)),(b,"الساعات",plan["total_hours"]),(c,"المستويات",8)]:
        with col: st.markdown(f'<div class="card"><div>{t}</div><div class="metric">{v}</div></div>',unsafe_allow_html=True)
elif st.session_state.page=="college":
    st.markdown('<div class="title">كلية اللغات وعلومها</div>',unsafe_allow_html=True); st.write("")
    st.markdown('<div class="card"><h3>اللغة الإنجليزية والترجمة</h3><p>برنامج بكالوريوس متكامل مع المقررات والملفات والاختبارات الذكية.</p></div>',unsafe_allow_html=True)
    if st.button("دخول البرنامج",use_container_width=True): st.session_state.page="program";st.rerun()
elif st.session_state.page=="program":
    st.markdown('<div class="title">بكالوريوس اللغة الإنجليزية والترجمة</div><div>الخطة الدراسية 2026 • 137 ساعة</div>',unsafe_allow_html=True)
    for level in range(1,9):
        st.markdown(f"### المستوى {level}"); items=[x for x in COURSES if x["level"]==level]; cols=st.columns(3)
        for j,c in enumerate(items):
            with cols[j%3]:
                st.markdown(f'<div class="card"><span class="badge">{c["code"]} • {c["hours"]} ساعات</span><h3>{c["name"]}</h3></div>',unsafe_allow_html=True)
                if st.button("فتح المقرر",key="o"+c["id"],use_container_width=True): st.session_state.cid=c["id"];st.session_state.page="course";st.rerun()
else:
    c=next(x for x in COURSES if x["id"]==st.session_state.cid)
    st.markdown(f'<div class="title">{c["name"]}</div><div>{c["code"]} • {c["hours"]} ساعات • المستوى {c["level"]}</div>',unsafe_allow_html=True)
    t1,t2,t3,t4=st.tabs(["📘 المقرر","📁 الملفات","🤖 الاختبار الذكي","💡 Pro-Tips"])
    with t1: st.markdown(f'<div class="card"><h3>عن المقرر</h3><p>{c["description"]}</p></div>',unsafe_allow_html=True)
    with t2:
        up=st.file_uploader("ارفع ملفات المقرر",type=["pdf","docx","pptx","txt","md"],accept_multiple_files=True)
        if up:
            for f in up: (folder(c)/f.name).write_bytes(f.getbuffer())
            st.success("تم حفظ الملفات.")
        for p in folder(c).glob("*"): st.write("📄",p.name)
        if not list(folder(c).glob("*")): st.info("ارفع ملفات المقرر هنا.")
    with t3:
        n=st.select_slider("عدد الأسئلة",[5,10,15],5)
        if st.button("✨ إنشاء اختبار ذكي",use_container_width=True):
            st.session_state.quiz,msg=make_quiz(c,provider,key,n);st.session_state.msg=msg;st.session_state.i=0;st.session_state.score=0;st.session_state.answered=False;st.rerun()
        if st.session_state.quiz:
            q=st.session_state.quiz; i=st.session_state.i
            if i<len(q):
                item=q[i]; st.info(st.session_state.msg); st.markdown(f'<div class="q"><b>السؤال {i+1} من {len(q)}</b><br><br>{item["question"]}</div>',unsafe_allow_html=True)
                if not st.session_state.answered:
                    ans=st.radio("الإجابة",item["options"],key=f"a{i}")
                    if st.button("تأكيد الإجابة",key=f"c{i}",use_container_width=True):
                        st.session_state.last=ans;st.session_state.answered=True
                        if ans==item["answer"]: st.session_state.score+=1
                        st.rerun()
                else:
                    if st.session_state.last==item["answer"]: st.markdown(f'<div class="ok"><b>✓ إجابة صحيحة</b><br>{item["explanation"]}</div>',unsafe_allow_html=True)
                    else: st.markdown(f'<div class="no"><b>✗ إجابة غير صحيحة</b><br>الصحيح: {item["answer"]}<br>{item["explanation"]}</div>',unsafe_allow_html=True)
                    if st.button("السؤال التالي →",key=f"n{i}",use_container_width=True): st.session_state.i+=1;st.session_state.answered=False;st.rerun()
            else:
                st.success(f"انتهى الاختبار: {st.session_state.score}/{len(q)} — {round(st.session_state.score/len(q)*100)}%")
    with t4:
        for tip in c["tips"]: st.markdown("• "+tip)

st.markdown('<div style="text-align:center;color:#64748b;padding:35px;font-size:12px">نموذج مستقل للعرض والتطوير، مستند إلى معلومات الخطة المنشورة رسميًا، ولا يمثل بوابة رسمية أو اعتمادًا من الجامعة.</div>',unsafe_allow_html=True)
